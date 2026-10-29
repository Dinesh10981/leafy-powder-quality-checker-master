"""Excel database schema initialization and migration utilities."""
from __future__ import annotations

from typing import Any, Callable
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..models.sensor_record import BANDS, COLUMN_WIDTHS, HEADERS


def apply_schema_migration(sheet: Any, analysis_fn: Callable[[dict[str, float]], dict[str, Any]] | None = None) -> None:
    """Ensure workbook schema, formatting, headers, table styles, and freeze panes."""
    for column, header in enumerate(HEADERS, start=1):
        sheet.cell(1, column, header)

    header_fill = PatternFill("solid", fgColor="006B2C")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:N{max(1, sheet.max_row)}"
    sheet.row_dimensions[1].height = 30

    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        col_letter = openpyxl.utils.get_column_letter(index)
        sheet.column_dimensions[col_letter].width = width

    for row in range(2, sheet.max_row + 1):
        if all(isinstance(sheet.cell(row, col).value, (int, float)) for col in range(4, 10)):
            values = {band: float(sheet.cell(row, col).value) for band, col in zip(BANDS, range(4, 10))}
            if analysis_fn:
                analysis = analysis_fn(values)
                sheet.cell(row, 10, analysis["red_edge_average"])
                sheet.cell(row, 11, analysis["nir_average"])
                sheet.cell(row, 12, analysis["lpqi"])
                sheet.cell(row, 13, analysis["quality"])
            if not sheet.cell(row, 14).value:
                sheet.cell(row, 14, "Imported reading")

        for col in range(4, 12):
            sheet.cell(row, col).number_format = "0.0000"
        sheet.cell(row, 12).number_format = "0.000"

    if not sheet.tables:
        table = Table(displayName="SensorReadings", ref=f"A1:N{max(2, sheet.max_row)}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    else:
        for table in sheet.tables.values():
            table.ref = f"A1:N{max(2, sheet.max_row)}"

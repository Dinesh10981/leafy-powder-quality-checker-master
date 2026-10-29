"""Thread-safe persistent ExcelStore database connection and repository."""
from __future__ import annotations

import threading
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Mapping

import openpyxl
from openpyxl import Workbook

from ..migrations.schema_migration import apply_schema_migration
from ..models.sensor_record import BANDS
from ..queries.sensor_queries import (
    count_records,
    delete_all_records,
    fetch_history,
    get_next_sample_number,
    row_to_dict,
)


class DatabaseStorageError(RuntimeError):
    """Raised when persistent storage fails or file is locked."""


class ExcelStore:
    """Thread-safe append/read access to the persistent Excel database."""

    def __init__(self, path: Path, analysis_fn: Callable[[dict[str, float]], dict[str, Any]] | None = None) -> None:
        self.path = Path(path).resolve()
        self.analysis_fn = analysis_fn
        self.lock = threading.RLock()
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self.lock:
            self.workbook, self.sheet = self._load_or_create()
            self._ensure_schema()
            self._save()

    def _load_or_create(self):
        if self.path.exists():
            workbook = openpyxl.load_workbook(self.path)
        else:
            workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sensor Readings"
        return workbook, sheet

    def _ensure_schema(self) -> None:
        apply_schema_migration(self.sheet, analysis_fn=self.analysis_fn)

    def _save(self) -> None:
        try:
            self.workbook.save(self.path)
        except PermissionError as exc:
            raise DatabaseStorageError(
                f"Cannot write {self.path.name}. Close it in Excel and scan again."
            ) from exc

    def append(
        self,
        values: Mapping[str, float],
        analysis: Mapping[str, float | str],
        sensor_url: str,
    ) -> dict[str, Any]:
        with self.lock:
            now = datetime.now()
            sample_number = get_next_sample_number(self.sheet)
            self.sheet.append([
                sample_number,
                now.strftime("%Y-%m-%d"),
                now.strftime("%H:%M:%S"),
                *(values[band] for band in BANDS),
                analysis["red_edge_average"],
                analysis["nir_average"],
                analysis["lpqi"],
                analysis["quality"],
                sensor_url,
            ])
            row = self.sheet.max_row
            for col in range(4, 12):
                self.sheet.cell(row, col).number_format = "0.0000"
            self.sheet.cell(row, 12).number_format = "0.000"
            self.sheet.auto_filter.ref = f"A1:N{row}"
            for table in self.sheet.tables.values():
                table.ref = f"A1:N{row}"
            self._save()
            record = row_to_dict(self.sheet, row)
            if "ml_prediction" in analysis:
                record["ml_prediction"] = analysis["ml_prediction"]
            if "confidence" in analysis:
                record["confidence"] = analysis["confidence"]
            return record

    def history(self, limit: int = 500) -> list[dict[str, Any]]:
        with self.lock:
            return fetch_history(self.sheet, limit=limit)

    def count(self) -> int:
        with self.lock:
            return count_records(self.sheet)

    def export_bytes(self) -> bytes:
        with self.lock:
            return self.path.read_bytes()

    def clear(self) -> None:
        with self.lock:
            delete_all_records(self.sheet)
            self._save()
